import csv
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).parent
SCRIPT = ROOT / "lab_data.py"


class LabDataCliTests(unittest.TestCase):
    def run_cli(self, *args, input_text=None):
        return subprocess.run(
            [sys.executable, str(SCRIPT), *map(str, args)],
            input=input_text,
            text=True,
            capture_output=True,
        )

    def test_generates_three_sheets_and_preserves_text_and_statistics(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            source = directory / "input.csv"
            output = directory / "report.xlsx"
            source.write_text(
                "\ufeffid,temp,ph,note\n"
                '001," 1.5 ",7,"quoted, note"\n'
                '002,,bad,"=1+1"\n'
                '003,NaN,9," leading text "\n',
                encoding="utf-8",
                newline="",
            )
            result = self.run_cli(source, "--columns", "temp,ph", "--output", output)
            self.assertEqual(result.returncode, 0, result.stderr)

            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook.sheetnames, ["原始", "汇总", "异常"])
            raw = workbook["原始"]
            self.assertEqual(list(raw.values), [
                ("id", "temp", "ph", "note"),
                ("001", " 1.5 ", "7", "quoted, note"),
                ("002", None, "bad", "=1+1"),
                ("003", "NaN", "9", " leading text "),
            ])
            self.assertEqual(raw["A2"].data_type, "s")
            self.assertEqual(raw["D3"].data_type, "s")
            self.assertEqual(list(workbook["汇总"].values), [
                ("column", "count", "missing", "invalid", "min", "max", "mean"),
                ("temp", 1, 1, 1, 1.5, 1.5, 1.5),
                ("ph", 2, 0, 1, 7, 9, 8),
            ])
            self.assertEqual(list(workbook["异常"].values), [
                ("data_row", "column", "value", "reason"),
                (2, "temp", None, "missing"),
                (3, "temp", "NaN", "invalid"),
                (2, "ph", "bad", "invalid"),
            ])

    def test_cli_preserves_missing_spaces_and_formula_text_in_derived_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            source = directory / "formula-input.csv"
            output = directory / "formula-report.xlsx"
            source.write_text(
                "=formula,note\n"
                "   ,plain\n"
                "=1+1,plain\n",
                encoding="utf-8",
                newline="",
            )
            result = self.run_cli(
                source, "--columns", "=formula", "--output", output
            )
            self.assertEqual(result.returncode, 0, result.stderr)

            workbook = load_workbook(output, data_only=False)
            summary = workbook["汇总"]
            issues = workbook["异常"]
            self.assertEqual(summary["A2"].value, "=formula")
            self.assertEqual(summary["A2"].data_type, "s")
            self.assertEqual(summary["B2"].value, 0)
            self.assertEqual(summary["C2"].value, 1)
            self.assertEqual(summary["D2"].value, 1)
            self.assertEqual(issues["B2"].value, "=formula")
            self.assertEqual(issues["B2"].data_type, "s")
            self.assertEqual(issues["C2"].value, "   ")
            self.assertEqual(issues["C2"].data_type, "s")
            self.assertEqual(issues["D2"].value, "missing")
            self.assertEqual(issues["C3"].value, "=1+1")
            self.assertEqual(issues["C3"].data_type, "s")
            self.assertEqual(issues["D3"].value, "invalid")

    def test_interactive_mode_matches_cli_and_existing_output_is_not_overwritten(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            source = directory / "input.csv"
            output = directory / "report.xlsx"
            source.write_text("id,value\na,2\n", encoding="utf-8")
            prompts = f"{source}\nvalue\n{output}\n"
            interactive = self.run_cli(input_text=prompts)
            self.assertEqual(interactive.returncode, 0, interactive.stderr)
            original = output.read_bytes()
            refused = self.run_cli(source, "--columns", "value", "--output", output)
            self.assertNotEqual(refused.returncode, 0)
            self.assertEqual(output.read_bytes(), original)

    def test_structural_errors_do_not_create_output(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            for name, contents in {
                "empty.csv": "",
                "duplicate.csv": "a,a\n1,2\n",
                "short.csv": "a,b\n1\n",
            }.items():
                source = directory / name
                output = directory / f"{name}.xlsx"
                source.write_text(contents, encoding="utf-8")
                result = self.run_cli(source, "--columns", "a", "--output", output)
                self.assertNotEqual(result.returncode, 0, name)
                self.assertFalse(output.exists(), name)

    def test_input_cannot_be_output_and_no_valid_values_are_blank(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            source = directory / "input.csv"
            source.write_text("x\nNaN\ninf\n \n", encoding="utf-8")
            refused = self.run_cli(source, "--columns", "x", "--output", source)
            self.assertNotEqual(refused.returncode, 0)
            workbook_path = directory / "out.xlsx"
            result = self.run_cli(source, "--columns", "x", "--output", workbook_path)
            self.assertEqual(result.returncode, 0, result.stderr)
            row = list(load_workbook(workbook_path)["汇总"].values)[1]
            self.assertEqual(row, ("x", 0, 1, 2, None, None, None))


if __name__ == "__main__":
    unittest.main()
